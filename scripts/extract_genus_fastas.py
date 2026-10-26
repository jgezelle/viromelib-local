#!/usr/bin/env python3
"""
Extract sequences from the annotated peaks spreadsheet into per-genus FASTA
files, keeping only rows that:
  - have n_reps_passing of 3 or 4 (i.e. column 'n_reps_passing' in {3, 4})
  - do NOT match any covariance model (column 'any_hit' == False)

Usage:
    python3 extract_genus_fastas.py <input.xlsx> <output_dir>

Requires: openpyxl  (pip install openpyxl)
"""

import sys
import os
import re
from collections import defaultdict

import openpyxl


def safe_filename(s):
    """Turn a genus name into a filesystem-safe filename."""
    return re.sub(r'[^A-Za-z0-9._-]', '_', str(s))


def main():
    if len(sys.argv) != 3:
        print(f"Usage: {sys.argv[0]} <input.xlsx> <output_dir>")
        sys.exit(1)

    in_path = sys.argv[1]
    out_dir = sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)

    wb = openpyxl.load_workbook(in_path, data_only=True)
    ws = wb[wb.sheetnames[0]]  # assumes the data is on the first sheet

    # Read header row to find column indices by name (robust to column reordering)
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col = {name: idx for idx, name in enumerate(header)}

    required = ['genus', 'species', 'peak_id', 'family',
                'n_reps_passing', 'downstream_seq', 'any_hit']
    missing = [c for c in required if c not in col]
    if missing:
        raise ValueError(f"Missing expected column(s) in sheet: {missing}")

    by_genus = defaultdict(list)
    total_qualifying = 0
    skipped_no_seq = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        n_reps_passing = row[col['n_reps_passing']]
        any_hit = row[col['any_hit']]

        if n_reps_passing in (3, 4) and any_hit is False:
            total_qualifying += 1
            seq = row[col['downstream_seq']]
            if not seq:
                skipped_no_seq += 1
                continue

            genus = row[col['genus']]
            by_genus[genus].append({
                'peak_id': row[col['peak_id']],
                'seq': seq,
                'n_reps_passing': n_reps_passing,
                'family': row[col['family']],
                'species': row[col['species']],
            })

    print(f"Qualifying rows (3 or 4 passing reps, no CM hit): {total_qualifying}")
    print(f"  Skipped (no sequence available): {skipped_no_seq}")
    print()

    written_files = []
    for genus, records in by_genus.items():
        fname = os.path.join(out_dir, f"{safe_filename(genus)}.fasta")
        with open(fname, 'w') as f:
            for rec in records:
                header_line = (
                    f">{rec['peak_id']} genus={genus} "
                    f"species={rec['species']} n_reps_passing={rec['n_reps_passing']}"
                )
                f.write(header_line + "\n")
                f.write(str(rec['seq']) + "\n")
        written_files.append(fname)
        print(f"  {genus}: {len(records)} sequences -> {fname}")

    # Also write a combined FASTA with everything
    combined_path = os.path.join(out_dir, "all_genera_combined.fasta")
    with open(combined_path, 'w') as out:
        for fname in written_files:
            with open(fname) as f:
                out.write(f.read())
    print(f"\nCombined FASTA written to: {combined_path}")


if __name__ == "__main__":
    main()
